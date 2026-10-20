# Report text extraction and model analysis service.
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from typing import Any

MAX_EXTRACTED_CHARS = 60000
MAX_OCR_PDF_PAGES = 40


class ReportExtractionError(RuntimeError):
    pass


def _clean_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()[:MAX_EXTRACTED_CHARS]


def _ocr_image(image: Any) -> str:
    try:
        import pytesseract
        from PIL import Image  # noqa: F401
    except ImportError as exc:
        raise ReportExtractionError(
            "图片 OCR 依赖未安装。请安装 Pillow、pytesseract 和系统 Tesseract。"
        ) from exc
    try:
        languages = pytesseract.get_languages(config="")
    except Exception as exc:
        raise ReportExtractionError(
            "无法启动 Tesseract OCR。请确认系统已安装 tesseract-ocr。"
        ) from exc
    if "chi_sim" not in languages:
        raise ReportExtractionError(
            "系统缺少中文 OCR 语言包 chi_sim。Ubuntu 可安装 tesseract-ocr-chi-sim 后重试。"
        )
    try:
        return pytesseract.image_to_string(image, lang="chi_sim+eng", timeout=45)
    except Exception as exc:
        raise ReportExtractionError("图片 OCR 失败，请检查图片清晰度或 OCR 安装情况。") from exc


def _extract_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = reader.pages
        text = "\n\n".join(
            (page.extract_text() or "") for page in pages[:MAX_OCR_PDF_PAGES]
        )
    except ImportError as exc:
        raise ReportExtractionError("缺少 PDF 提取依赖 pypdf，请按安装说明安装新增依赖。") from exc
    except Exception as exc:
        raise ReportExtractionError(
            "无法读取该 PDF。请确认文件未损坏且未设置无法访问的密码。"
        ) from exc

    if len(text.strip()) >= 40:
        if len(pages) > MAX_OCR_PDF_PAGES:
            text += f"\n\n[提示：文档共 {len(pages)} 页，本次最多读取前 {MAX_OCR_PDF_PAGES} 页。]"
        return _clean_text(text), "pdf_text"

    try:
        import fitz
        from PIL import Image
    except ImportError as exc:
        raise ReportExtractionError(
            "该 PDF 可能是扫描件。请安装 PyMuPDF、Pillow、pytesseract 和系统中文 Tesseract OCR 后重试。"
        ) from exc

    try:
        document = fitz.open(str(path))
        if document.page_count == 0:
            raise ReportExtractionError("PDF 没有可读取的页面。")
        pieces: list[str] = []
        for index in range(min(document.page_count, MAX_OCR_PDF_PAGES)):
            page = document.load_page(index)
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image = Image.open(BytesIO(pix.tobytes("png")))
            pieces.append(_ocr_image(image))
        if document.page_count > MAX_OCR_PDF_PAGES:
            pieces.append(
                f"[提示：文档共 {document.page_count} 页，本次最多识别前 {MAX_OCR_PDF_PAGES} 页。]"
            )
        result = _clean_text("\n\n".join(pieces))
        if len(result.strip()) < 10:
            raise ReportExtractionError("未能从扫描版 PDF 识别出足够文本，请上传更清晰的文件。")
        return result, "pdf_ocr"
    except ReportExtractionError:
        raise
    except Exception as exc:
        raise ReportExtractionError("扫描版 PDF OCR 失败，请检查文件和中文 OCR 环境。") from exc


def _extract_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document
        document = Document(str(path))
    except ImportError as exc:
        raise ReportExtractionError("缺少 DOCX 提取依赖 python-docx，请按安装说明安装新增依赖。") from exc
    except Exception as exc:
        raise ReportExtractionError("无法读取该 DOCX 文件，文件可能已损坏。") from exc

    pieces = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            pieces.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = _clean_text("\n".join(pieces))
    if not text:
        raise ReportExtractionError("DOCX 中没有提取到正文文本。")
    return text, "docx_text"


def extract_report_text(path: Path, filename: str) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text, method = _extract_pdf(path)
    elif suffix == ".docx":
        text, method = _extract_docx(path)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        try:
            from PIL import Image
            with Image.open(path) as image:
                text = _clean_text(_ocr_image(image.convert("RGB")))
        except ReportExtractionError:
            raise
        except ImportError as exc:
            raise ReportExtractionError("缺少图片 OCR 依赖 Pillow，请按安装说明安装新增依赖。") from exc
        except Exception as exc:
            raise ReportExtractionError("无法读取该图片或 OCR 识别失败。") from exc
        method = "image_ocr"
    else:
        raise ReportExtractionError("当前文件格式不支持分析。")
    if len(text.strip()) < 10:
        raise ReportExtractionError("报告中没有提取到足够的可读文本。请上传清晰的 PDF 或图片。")
    return _clean_text(text), method


def generate_report_analysis(text: str, filename: str) -> tuple[str, str]:
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
        from langchain_openai import ChatOpenAI
        from agent.config.settings import get_settings

        config = get_settings().get_model_config()
        model_name = str(config.get("model") or config.get("model_name") or "configured-model")
        llm = ChatOpenAI(**config, temperature=0)
        system_prompt = (
            "你是心理测评报告的结构化解读助手，不是临床诊断系统。\n"
            "1. 仅依据报告文本，不补造分数、量表版本、常模、引用或结论。\n"
            "2. 区分报告明确记载的事实、可能的解释和无法确认的信息。\n"
            "3. 量表版本、评分规则、答案或常模不完整时，不自行推算或作确定性严重程度判断。\n"
            "4. 不作精神障碍诊断，不开药，不给个体化药物剂量，不替代专业评估。\n"
            "5. 报告若提及自伤/自杀内容，指出需专业人员及时评估；若提示当前迫切危险，建议立即联系当地急救、前往急诊并联系可信任的人陪伴。\n"
            "6. OCR 可能出错，标出疑似识别错误、缺失字段和需人工核对内容。\n"
            "7. 使用中文 Markdown，采用以下标题：\n"
            "## 报告基本信息\n## 关键结果摘录\n## 结果解读与适用边界\n"
            "## 建议进一步核对\n## 求助提示\n"
            "报告未提供的信息请写“报告未提供”，不要猜测。"
        )
        user_prompt = (
            f"文件名：{filename}\n\n"
            "请分析以下心理测评报告文本。分数与量表名称请忠实摘录；"
            "若文本可能存在 OCR 错误，请明确提示人工核对。\n\n"
            f"<报告文本>\n{text[:MAX_EXTRACTED_CHARS]}\n</报告文本>"
        )
        response = llm.invoke([
            SystemMessage(content=system_prompt),
            HumanMessage(content=user_prompt),
        ])
        content = response.content
        if isinstance(content, list):
            content = "\n".join(
                str(item.get("text", "")) if isinstance(item, dict) else str(item)
                for item in content
            )
        result = str(content).strip()
        if not result:
            raise RuntimeError("模型返回了空内容")
        return result, model_name
    except Exception as exc:
        raise RuntimeError(
            "模型分析失败。请检查后端模型配置、代理和服务日志后重试。"
        ) from exc
